"""
modules/resume_parser.py — Extract raw text from PDF and DOCX resume files.

This module is responsible ONLY for text extraction.
Skill identification is handled by skill_extractor.py.

Supported formats:
  - PDF  (via pdfplumber; falls back to a user-readable error for scanned PDFs)
  - DOCX (via python-docx)

Design notes:
  - Returns a dict with keys: text, pages, warnings, success.
  - Never raises exceptions to callers; surfaces errors in the warnings list.
  - File-size and extension validation is the caller's responsibility
    (enforced in app.py before calling these functions).
"""

from __future__ import annotations

import io
import os
import re
from typing import Dict, List, Optional


# ---------------------------------------------------------------------------
# PDF parsing
# ---------------------------------------------------------------------------

def parse_pdf(file_bytes: bytes, filename: str = "resume.pdf") -> Dict:
    """
    Extract text from a PDF given its raw bytes.

    Returns:
        {
          "text": str,          # full extracted text
          "pages": int,         # number of pages
          "warnings": list[str],
          "success": bool
        }
    """
    warnings: List[str] = []
    text_parts: List[str] = []
    pages = 0

    try:
        import pdfplumber  # type: ignore
    except ImportError:
        return {
            "text": "",
            "pages": 0,
            "warnings": ["pdfplumber is not installed. Run: pip install pdfplumber"],
            "success": False,
        }

    try:
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            pages = len(pdf.pages)
            for i, page in enumerate(pdf.pages, start=1):
                try:
                    page_text = page.extract_text()
                    if page_text:
                        text_parts.append(page_text)
                    else:
                        warnings.append(
                            f"Page {i} returned no extractable text. "
                            "It may contain scanned images or non-standard encoding."
                        )
                except Exception as page_err:  # noqa: BLE001
                    warnings.append(f"Page {i} could not be read: {page_err}")
    except Exception as err:  # noqa: BLE001
        return {
            "text": "",
            "pages": pages,
            "warnings": [f"Failed to open PDF '{filename}': {err}"],
            "success": False,
        }

    full_text = "\n".join(text_parts)
    full_text = _clean_text(full_text)

    if not full_text.strip():
        warnings.append(
            "No text could be extracted from this PDF. "
            "This usually means it is a scanned image. "
            "Please upload a text-based PDF or manually enter your skills."
        )
        return {"text": "", "pages": pages, "warnings": warnings, "success": False}

    return {"text": full_text, "pages": pages, "warnings": warnings, "success": True}


# ---------------------------------------------------------------------------
# DOCX parsing
# ---------------------------------------------------------------------------

def parse_docx(file_bytes: bytes, filename: str = "resume.docx") -> Dict:
    """
    Extract text from a DOCX file given its raw bytes.

    Extracts from:
      - Body paragraphs (including heading styles)
      - Table cells
      - Text boxes are not accessible via python-docx; a warning is added.
    """
    warnings: List[str] = []

    try:
        from docx import Document  # type: ignore
    except ImportError:
        return {
            "text": "",
            "pages": 0,
            "warnings": ["python-docx is not installed. Run: pip install python-docx"],
            "success": False,
        }

    try:
        doc = Document(io.BytesIO(file_bytes))
    except Exception as err:  # noqa: BLE001
        return {
            "text": "",
            "pages": 0,
            "warnings": [f"Failed to open DOCX '{filename}': {err}"],
            "success": False,
        }

    text_parts: List[str] = []

    # Body paragraphs
    for para in doc.paragraphs:
        if para.text.strip():
            text_parts.append(para.text)

    # Table cells
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text_parts.append(cell.text)

    full_text = "\n".join(text_parts)
    full_text = _clean_text(full_text)

    warnings.append(
        "Note: Text boxes embedded in DOCX files cannot be read automatically. "
        "If your resume uses text boxes, please add those skills manually."
    )

    if not full_text.strip():
        return {
            "text": "",
            "pages": 0,
            "warnings": [f"No text could be extracted from '{filename}'."] + warnings,
            "success": False,
        }

    return {"text": full_text, "pages": 0, "warnings": warnings, "success": True}


# ---------------------------------------------------------------------------
# Dispatcher
# ---------------------------------------------------------------------------

def parse_resume(file_bytes: bytes, filename: str) -> Dict:
    """
    Route to the correct parser based on file extension.

    Adds an 'extension' key to the result dict.
    """
    ext = os.path.splitext(filename.lower())[1]
    if ext == ".pdf":
        result = parse_pdf(file_bytes, filename)
    elif ext == ".docx":
        result = parse_docx(file_bytes, filename)
    else:
        result = {
            "text": "",
            "pages": 0,
            "warnings": [
                f"Unsupported file format '{ext}'. "
                "Please upload a .pdf or .docx file."
            ],
            "success": False,
        }
    result["extension"] = ext
    result["filename"] = filename
    return result


# ---------------------------------------------------------------------------
# Text cleaning
# ---------------------------------------------------------------------------

def _clean_text(text: str) -> str:
    """
    Normalise whitespace and remove common PDF extraction artefacts.
    Does NOT alter content or case — case-normalisation happens in skill_extractor.
    """
    # Replace non-breaking spaces and other Unicode whitespace
    text = text.replace("\xa0", " ").replace("\u200b", "")
    # Collapse runs of whitespace within a line (but preserve newlines)
    lines = text.splitlines()
    lines = [re.sub(r" {2,}", " ", line).strip() for line in lines]
    # Remove blank lines in excess (keep single blank lines as section separators)
    cleaned_lines: List[str] = []
    prev_blank = False
    for line in lines:
        if line == "":
            if not prev_blank:
                cleaned_lines.append(line)
            prev_blank = True
        else:
            cleaned_lines.append(line)
            prev_blank = False
    return "\n".join(cleaned_lines).strip()


# ---------------------------------------------------------------------------
# Section detection (heuristic, used by skill_extractor)
# ---------------------------------------------------------------------------

_SECTION_HEADERS = {
    "skills": [
        "skills", "technical skills", "key skills", "core competencies",
        "technologies", "tech stack", "tools", "languages", "frameworks",
        "expertise", "competencies", "proficiencies",
    ],
    "education": [
        "education", "academic background", "academic qualifications",
        "qualifications", "degrees",
    ],
    "experience": [
        "experience", "work experience", "professional experience",
        "employment history", "internships", "internship", "work history",
    ],
    "projects": [
        "projects", "personal projects", "academic projects",
        "portfolio", "notable projects",
    ],
    "certifications": [
        "certifications", "certificates", "courses", "training",
        "professional development", "licences", "credentials",
    ],
    "summary": [
        "summary", "profile", "objective", "about me", "career objective",
        "professional summary",
    ],
}

# Compiled pattern: line is a section header if it matches a known phrase
# and is short (under ~60 chars) with no sentence-ending punctuation.
_HEADER_RE = re.compile(
    r"^\s*(" + "|".join(
        re.escape(h)
        for headers in _SECTION_HEADERS.values()
        for h in headers
    ) + r")\s*[:\-]?\s*$",
    re.IGNORECASE,
)


def detect_sections(text: str) -> Dict[str, str]:
    """
    Split resume text into named sections.
    Returns dict mapping section name → section text.
    Unknown lines are placed in 'other'.
    """
    sections: Dict[str, List[str]] = {"other": []}
    current_section = "other"
    current_section_name = "other"

    for line in text.splitlines():
        # Check if this line is a section header
        matched_section = _match_header(line)
        if matched_section:
            current_section = matched_section
            if current_section not in sections:
                sections[current_section] = []
        else:
            sections.setdefault(current_section, []).append(line)

    return {k: "\n".join(v).strip() for k, v in sections.items() if v}


def _match_header(line: str) -> Optional[str]:
    """Return the canonical section name if line is a section header, else None."""
    clean = line.strip().rstrip(":- ").lower()
    if not clean or len(clean) > 60:
        return None
    for section_name, headers in _SECTION_HEADERS.items():
        if clean in headers:
            return section_name
    return None
